"""bookfig – einheitliche Buchabbildungen als editierbare PPTX + PDF + PNG.

Alle Maße in mm. Eine Abbildung = eine Folie in exakter Endgröße.
Ein Theme ordnet semantischen Rollen (box, box_hi, box_strong, ...) Farben,
Linien und Schrift zu; die Abbildungsskripte verwenden nur Rollen.
"""
from __future__ import annotations

import math
import os
import re
import subprocess
import shutil
from dataclasses import dataclass, field

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

MM = 36000  # EMU je mm
PT_MM = 25.4 / 72

def _find_font(bold: bool, italic: bool) -> str:
    """Arial-metrische Schrift für die Textmessung: Arial, sonst Liberation Sans."""
    style = {(False, False): "Regular", (True, False): "Bold", (False, True): "Italic",
             (True, True): "Bold Italic"}[(bold, italic)]
    for fam in ("Arial", "Liberation Sans"):
        try:
            out = subprocess.run(["fc-match", "-f", "%{file}|%{family}", f"{fam}:style={style}"],
                                 capture_output=True, text=True, timeout=10).stdout
            path, family = out.split("|", 1)
            if fam.lower() in family.lower() and os.path.exists(path):
                return path
        except Exception:
            pass
    fallback = {(False, False): "LiberationSans-Regular.ttf", (True, False): "LiberationSans-Bold.ttf",
                (False, True): "LiberationSans-Italic.ttf", (True, True): "LiberationSans-BoldItalic.ttf"}
    return os.path.join("/usr/share/fonts/truetype/liberation", fallback[(bold, italic)])


_FONTS = {(b, i): ImageFont.truetype(_find_font(b, i), 1000) for b in (False, True) for i in (False, True)}

SOFFICE_WRAPPER = "/mnt/skills/public/pptx/scripts/office/soffice.py"


def pptx_to_pdf(pptx_path: str, outdir: str):
    """PPTX -> PDF über LibreOffice. Nutzt den Sandbox-Wrapper des pptx-Skills, falls vorhanden."""
    if os.path.exists(SOFFICE_WRAPPER):
        cmd = ["python3", SOFFICE_WRAPPER, "--headless", "--convert-to", "pdf", "--outdir", outdir, pptx_path]
        subprocess.run(cmd, check=True, capture_output=True, timeout=240)
        return
    import tempfile
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as prof:
        env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
        cmd = ["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf",
               "--outdir", outdir, pptx_path]
        subprocess.run(cmd, check=True, capture_output=True, timeout=240, env=env)


# --------------------------------------------------------------------------
# Farbe / Kontrast
# --------------------------------------------------------------------------
def _lin(c):
    c = c / 255
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def luminance(hexcol: str) -> float:
    r, g, b = int(hexcol[0:2], 16), int(hexcol[2:4], 16), int(hexcol[4:6], 16)
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def contrast(a: str, b: str) -> float:
    la, lb = sorted([luminance(a), luminance(b)], reverse=True)
    return (la + 0.05) / (lb + 0.05)


def gray_value(hexcol: str) -> int:
    """Graustufe 0..255 wie bei S/W-Druck (Rec. 601 Luma)."""
    r, g, b = int(hexcol[0:2], 16), int(hexcol[2:4], 16), int(hexcol[4:6], 16)
    return round(0.299 * r + 0.587 * g + 0.114 * b)


# --------------------------------------------------------------------------
# Theme
# --------------------------------------------------------------------------
@dataclass
class Role:
    fill: str | None = None        # Hex ohne '#', None = keine Füllung
    line: str | None = None        # Hex, None = keine Linie
    lw: float = 0.5                # Linienstärke in pt
    dash: str | None = None        # None | 'dash' | 'sysDot' | 'sysDash'
    text: str = "1A1A1A"           # Textfarbe auf dieser Fläche
    radius: float = 0.8            # Eckradius mm (0 = eckig)


@dataclass
class Theme:
    key: str
    name: str
    font: str = "Arial"
    sizes: dict = field(default_factory=lambda: {
        "body": 7.0, "small": 6.0, "label": 7.0, "head": 7.5, "note": 6.5, "band": 6.0})
    text: str = "1A1A1A"
    text_muted: str = "4D4D4D"
    roles: dict = field(default_factory=dict)
    arrow: str = "333F44"           # Standard-Verbinder
    arrow_lw: float = 0.75
    accent: str = "C62828"          # Klammern, Prozesspfeile
    info: str = "6B7A80"            # gestrichelte Informationsflüsse
    band_label: str = "4D4D4D"
    band_caps: bool = True

    def role(self, name: str) -> Role:
        return self.roles[name]


def make_themes() -> dict[str, Theme]:
    """Festgelegter Stil der 6. Auflage: Graustufen-Grundlage, Anthrazit für Struktur,
    Rot (#C62828) als einzige Kontrastfarbe für das jeweils wichtigste Element."""
    T = {}
    T["buch"] = Theme(
        key="buch", name="6. Auflage · Anthrazit + Rot",
        roles={
            "box":        Role(fill="EFEFEF", line="9A9A9A", lw=0.5),     # Standardbox
            "box_hi":     Role(fill="CFCFCF", line="8A8A8A", lw=0.5),     # hervorgehobene Kategorie (z. B. monetäre Pläne)
            "box_strong": Role(fill="C62828", line=None, text="FFFFFF"),  # Schlüsselelement, max. 1–2 je Abbildung
            "box_plain":  Role(fill="FFFFFF", line="9A9A9A", lw=0.5),     # weiße Zelle in einer Box
            "container":  Role(fill=None, line="8C8C8C", lw=0.5, dash="sysDash", radius=1.2),  # Gruppe
            "band":       Role(fill="F6F6F6", line=None, radius=1.5),     # Hintergrundband (Ebene/Phase)
            "header":     Role(fill="3A3F44", line=None, text="FFFFFF", radius=0.8),  # Kopfleiste
        },
        arrow="3A3F44", accent="C62828", info="6E6E6E", band_label="3A3F44",
    )
    return T


THEMES = make_themes()
THEME = THEMES["buch"]

# Gestaltungsbreite = Satzspiegelbreite der 5. Auflage; LaTeX skaliert das PNG auf \textwidth.
FIG_W = 110.0
FIG_W_QUER = 155.0   # ganzseitige Querabbildung (in LaTeX als PNG um 90° gedreht), Höhe max. 98 mm


# --------------------------------------------------------------------------
# Textmessung und Umbruch
# --------------------------------------------------------------------------
def text_width_mm(s: str, size_pt: float, bold=False, italic=False) -> float:
    f = _FONTS[(bold, italic)]
    return f.getlength(s) / 1000 * size_pt * PT_MM


SHY = chr(0xAD)  # weiches Trennzeichen (bewusst ohne Escape-Sequenz notiert)


def _segments(s: str):
    """Zerlegt in Stücke, nach denen umbrochen werden darf: Leerzeichen,
    '-' oder '/' vor Buchstaben, weiches Trennzeichen."""
    segs, cur = [], ""
    for i, ch in enumerate(s):
        if ch == SHY:
            segs.append((cur, "shy")); cur = ""; continue
        cur += ch
        nxt = s[i + 1] if i + 1 < len(s) else ""
        if ch == " ":
            segs.append((cur, "sp")); cur = ""
        elif ch in "-/" and nxt.isalpha() and i > 0 and (s[i - 1].isalpha() or s[i - 1] in "-/"):
            tail = s[i + 1:].split(" ")[0].rstrip(",.;:)")
            if len(tail) >= 6:          # keine Trennung vor kurzem Rest ("Mengen-Ziele")
                segs.append((cur, "hy")); cur = ""
    segs.append((cur, "end"))
    return [sg for sg in segs if sg[0] != "" or sg[1] == "end"]


def wrap(s: str, width_mm: float, size_pt: float, bold=False, italic=False) -> list[str]:
    """Greedy-Umbruch; harte Umbrüche mit '\n'. Liefert Zeilen ohne weiche Trennzeichen."""
    W = lambda t: text_width_mm(t, size_pt, bold, italic)
    out = []
    for hard in s.split("\n"):
        line, last = "", "end"
        for txt, kind in _segments(hard):
            trial = line + txt
            need = trial.rstrip() + ("-" if kind == "shy" else "")
            if line == "" or W(need) <= width_mm:
                line, last = trial, kind
            else:
                out.append(line.rstrip() + ("-" if last == "shy" else ""))
                line, last = txt, kind
        out.append(line.rstrip())
    return out


# --------------------------------------------------------------------------
# Absatz-Spezifikation
# --------------------------------------------------------------------------
@dataclass
class Para:
    text: str
    size: str = "body"         # Schlüssel aus theme.sizes oder Zahl
    bold: bool = False
    italic: bool = False
    color: str | None = None   # None = Rollenfarbe
    align: str | None = None   # None = Box-Default
    bullet: str | None = None  # z. B. "•"
    level: int = 0             # Einzug für Aufzählung
    space_before: float = 0.0  # pt
    nowrap: bool = False
    cont: bool = False         # Fortsetzung unter Aufzählungspunkt (gleicher Einzug, ohne Zeichen)


def P(text, **kw):
    return Para(text, **kw)


# --------------------------------------------------------------------------
# Figur
# --------------------------------------------------------------------------
class Fig:
    INS_X = 0.9   # mm Innenabstand links/rechts
    INS_Y = 0.6   # mm Innenabstand oben/unten
    LINE = 1.18   # Zeilenhöhe als Faktor der Schriftgröße

    def __init__(self, w_mm: float, h_mm: float, theme: Theme, name: str = "fig"):
        self.w, self.h, self.t, self.name = w_mm, h_mm, theme, name
        self.prs = Presentation()
        self.prs.slide_width = Emu(round(w_mm * MM))
        self.prs.slide_height = Emu(round(h_mm * MM))
        layout = self.prs.slide_layouts[6]
        self.slide = self.prs.slides.add_slide(layout)
        self.sh = self.slide.shapes
        self.warnings: list[str] = []
        self.text_checks: list[tuple[str, str, str]] = []  # (label, text color, bg)

    # ---------------- Hilfen ----------------
    def size(self, s):
        return self.t.sizes[s] if isinstance(s, str) else float(s)

    def measure(self, paras, w, ins=None):
        """Höhe (mm) eines Absatzblocks in einer Box der Breite w, inkl. Innenabstand."""
        ix, iy = ins if ins else (self.INS_X, self.INS_Y)
        inner_w = w - 2 * ix
        total = 0.0
        for para in paras:
            if isinstance(para, str):
                para = Para(para)
            sz = self.size(para.size)
            indent = (2.2 + para.level * 2.6) if (para.bullet or para.cont) else para.level * 2.6
            n = 1 if para.nowrap else len(wrap(para.text, inner_w - indent - 0.8, sz, para.bold, para.italic))
            total += n * sz * self.LINE * PT_MM + para.space_before * PT_MM
        return total + 2 * iy

    @staticmethod
    def _strip_style(shape):
        st = shape._element.find(qn("p:style"))
        if st is not None:
            shape._element.remove(st)

    def _apply_role(self, shape, role: Role, fill=None, line=None, lw=None, dash=None):
        fill = role.fill if fill is None else fill
        line = role.line if line is None else line
        lw = role.lw if lw is None else lw
        dash = role.dash if dash is None else dash
        if fill in (None, "none"):
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line in (None, "none"):
            shape.line.fill.background()
        else:
            shape.line.color.rgb = RGBColor.from_string(line)
            shape.line.width = Pt(lw)
            ln = shape._element.spPr.find(qn("a:ln"))
            if dash:
                pd = etree.SubElement(ln, qn("a:prstDash"))
                pd.set("val", dash)
            etree.SubElement(ln, qn("a:round"))

    def _set_radius(self, shape, w, h, r):
        if r <= 0:
            return
        adj = max(0, min(50000, round(r / min(w, h) * 100000)))
        shape.adjustments[0] = adj / 100000

    # ---------------- Text ----------------
    def _fill_text(self, shape, paras, w, h, role_text, align="c", anchor="m",
                   ins=None, label="", bg=None, autowrap=True):
        tf = shape.text_frame
        ix, iy = ins if ins else (self.INS_X, self.INS_Y)
        bp = tf._txBody.find(qn("a:bodyPr"))
        bp.set("lIns", str(round(ix * MM))); bp.set("rIns", str(round(ix * MM)))
        bp.set("tIns", str(round(iy * MM))); bp.set("bIns", str(round(iy * MM)))
        bp.set("wrap", "square")
        bp.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
        for child in list(bp):
            bp.remove(child)
        etree.SubElement(bp, qn("a:noAutofit"))
        inner_w = w - 2 * ix
        total_h = 0.0
        al_map = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
        first = True
        for para in paras:
            if isinstance(para, str):
                para = Para(para)
            sz = self.size(para.size)
            indent = (2.2 + para.level * 2.6) if (para.bullet or para.cont) else para.level * 2.6
            avail = inner_w - indent - 0.8
            lines = [para.text.replace(SHY, "")] if para.nowrap else wrap(para.text, avail, sz, para.bold, para.italic)
            for ln_ in lines:
                lw_ = text_width_mm(ln_, sz, para.bold, para.italic)
                if lw_ > avail + 0.05:
                    self.warnings.append(f"[{label}] Zeile zu breit ({lw_:.1f} > {avail:.1f} mm): {ln_!r}")
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = al_map[para.align or align]
            pPr = p._p.get_or_add_pPr()
            # Zeilenabstand fest in pt
            lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
            etree.SubElement(lnSpc, qn("a:spcPts")).set("val", str(round(sz * self.LINE * 100)))
            if para.space_before:
                sb = etree.SubElement(pPr, qn("a:spcBef"))
                etree.SubElement(sb, qn("a:spcPts")).set("val", str(round(para.space_before * 100)))
            if para.bullet:
                pPr.set("marL", str(round(indent * MM)))
                pPr.set("indent", str(round(-2.2 * MM)))
                bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", self.t.font)
                etree.SubElement(pPr, qn("a:buChar")).set("char", para.bullet)
            else:
                if indent:
                    pPr.set("marL", str(round(indent * MM)))
                etree.SubElement(pPr, qn("a:buNone"))
            color = para.color or role_text
            for i, ln_ in enumerate(lines):
                if i > 0:
                    br = etree.SubElement(p._p, qn("a:br"))
                    self._rpr(br, sz, para.bold, para.italic, color, tag="a:rPr")
                r = p.add_run()
                r.text = ln_
                self._style_run(r, sz, para.bold, para.italic, color)
            total_h += len(lines) * sz * self.LINE * PT_MM + para.space_before * PT_MM
            if bg:
                self.text_checks.append((label, color, bg))
        avail_h = h - 2 * iy
        if total_h > avail_h + 0.15:
            self.warnings.append(f"[{label}] Text zu hoch ({total_h:.1f} > {avail_h:.1f} mm)")
        return total_h

    def _rpr(self, parent, sz, bold, italic, color, tag="a:rPr"):
        rPr = etree.SubElement(parent, qn(tag))
        rPr.set("lang", "de-DE"); rPr.set("sz", str(round(sz * 100)))
        rPr.set("b", "1" if bold else "0"); rPr.set("i", "1" if italic else "0")
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(rPr, qn("a:latin")).set("typeface", self.t.font)
        etree.SubElement(rPr, qn("a:cs")).set("typeface", self.t.font)
        return rPr

    def _style_run(self, r, sz, bold, italic, color):
        rPr = r._r.get_or_add_rPr()
        rPr.set("lang", "de-DE")
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = self.t.font
        r.font.color.rgb = RGBColor.from_string(color)

    # ---------------- Formen ----------------
    def box(self, x, y, w, h, paras=None, role="box", align="c", anchor="m",
            name=None, rounded=True, fill=None, line=None, lw=None, dash=None,
            ins=None, rot=0, text_color=None):
        r = self.t.role(role)
        rad = r.radius if rounded else 0
        shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rad > 0 else MSO_SHAPE.RECTANGLE
        if rot in (90, 270, -90):
            # gedrehte Box: Form mit vertauschten Maßen anlegen und drehen
            cx, cy = x + w / 2, y + h / 2
            s = self.sh.add_shape(shp_type, Emu(round((cx - h / 2) * MM)), Emu(round((cy - w / 2) * MM)),
                                  Emu(round(h * MM)), Emu(round(w * MM)))
            s.rotation = 270 if rot in (270, -90) else 90
            bw, bh = h, w
        else:
            s = self.sh.add_shape(shp_type, Emu(round(x * MM)), Emu(round(y * MM)),
                                  Emu(round(w * MM)), Emu(round(h * MM)))
            bw, bh = w, h
        self._strip_style(s)
        if rad > 0:
            self._set_radius(s, bw, bh, rad)
        self._apply_role(s, r, fill, line, lw, dash)
        if name:
            s.name = name
        if paras:
            bg = fill if fill not in (None, "none") else (r.fill or "FFFFFF")
            self._fill_text(s, paras, bw, bh, text_color or r.text, align, anchor, ins,
                            label=name or (paras[0].text if isinstance(paras[0], Para) else paras[0])[:30], bg=bg)
        else:
            s.text_frame.text = ""
        return s

    def text(self, x, y, w, h, paras, align="l", anchor="t", color=None, bg="FFFFFF",
             name=None, ins=(0.0, 0.0), rot=0):
        if rot in (90, 270, -90):
            cx, cy = x + w / 2, y + h / 2
            s = self.sh.add_textbox(Emu(round((cx - h / 2) * MM)), Emu(round((cy - w / 2) * MM)),
                                    Emu(round(h * MM)), Emu(round(w * MM)))
            s.rotation = 270 if rot in (270, -90) else 90
            bw, bh = h, w
        else:
            s = self.sh.add_textbox(Emu(round(x * MM)), Emu(round(y * MM)), Emu(round(w * MM)), Emu(round(h * MM)))
            bw, bh = w, h
        if name:
            s.name = name
        lbl = name or (paras[0].text if isinstance(paras[0], Para) else paras[0])[:30]
        self._fill_text(s, paras, bw, bh, color or self.t.text, align, anchor, ins, label=lbl, bg=bg)
        return s

    def poly(self, pts, closed=True, fill=None, line=None, lw=0.5, dash=None,
             head=None, tail=None, name=None, arrow_size="med"):
        """Freiform aus Punkten (mm). head/tail: Pfeilspitze am Anfang/Ende."""
        fb = self.sh.build_freeform(round(pts[0][0] * MM), round(pts[0][1] * MM), scale=1.0)
        fb.add_line_segments([(round(px * MM), round(py * MM)) for px, py in pts[1:]], close=closed)
        s = fb.convert_to_shape()
        self._strip_style(s)
        if fill:
            s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(lw)
            ln = s._element.spPr.find(qn("a:ln"))
            if dash:
                etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
            etree.SubElement(ln, qn("a:round"))
            for tag, typ in (("a:headEnd", head), ("a:tailEnd", tail)):
                if typ:
                    e = etree.SubElement(ln, qn(tag))
                    e.set("type", typ); e.set("w", arrow_size); e.set("len", arrow_size)
        else:
            s.line.fill.background()
        if name:
            s.name = name
        return s

    def arrow(self, pts, role="arrow", head=False, tail=True, dash=None, lw=None, color=None,
              name=None, size="med"):
        """Verbinder als offene Polylinie (gerade oder mit Knicken)."""
        col = color or (self.t.info if role == "info" else self.t.accent if role == "accent" else self.t.arrow)
        lw = lw or (0.6 if role == "info" else self.t.arrow_lw)
        return self.poly(pts, closed=False, line=col, lw=lw, dash=dash,
                         head="triangle" if head else None, tail="triangle" if tail else None,
                         name=name, arrow_size=size)

    def line(self, pts, color=None, lw=None, dash=None, name=None):
        return self.arrow(pts, tail=False, head=False, color=color, lw=lw, dash=dash, name=name)

    def card(self, x, y, w, h, title, paras, head_h=5.0, name=None):
        """Karte: Kopfleiste (Rolle header, oben gerundet) + Rumpf (Rolle box)."""
        body = self.box(x, y, w, h, None, role="box", name=f"Karte {title}")
        hr = self.t.role("header")
        s = self.sh.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, Emu(round(x * MM)), Emu(round(y * MM)),
                              Emu(round(w * MM)), Emu(round(head_h * MM)))
        self._strip_style(s)
        r = self.t.role("box").radius
        s.adjustments[0] = r / head_h
        s.adjustments[1] = 0.0
        self._apply_role(s, hr)
        s.name = f"Kopf {title}"
        self._fill_text(s, [Para(title, size="head", bold=True)], w, head_h, hr.text, "l", "m",
                        (1.6, 0.3), label=title, bg=hr.fill)
        self.text(x, y + head_h, w, h - head_h, paras, align="l", anchor="t", ins=(1.6, 1.2),
                  bg=self.t.role("box").fill or "FFFFFF", name=f"Inhalt {title}")
        return body

    def ellipse(self, x, y, w, h, role="box", paras=None, fill=None, line=None, lw=None, name=None):
        r = self.t.role(role)
        s = self.sh.add_shape(MSO_SHAPE.OVAL, Emu(round(x * MM)), Emu(round(y * MM)),
                              Emu(round(w * MM)), Emu(round(h * MM)))
        self._strip_style(s)
        self._apply_role(s, r, fill, line, lw)
        if name:
            s.name = name
        if paras:
            self._fill_text(s, paras, w, h, r.text, "c", "m", (0.4, 0.4), label=name or "ellipse",
                            bg=fill or r.fill or "FFFFFF")
        return s

    def band(self, x, y, w, h, label, name=None):
        """Hintergrundband mit Beschriftung oben links."""
        self.box(x, y, w, h, None, role="band", name=name or f"Band {label}")
        txt = label.upper() if self.t.band_caps else label
        self.text(x + 1.6, y + 0.9, w - 3, 3.2,
                  [Para(txt, size="band", bold=True, color=self.t.band_label)],
                  name=f"Bandtitel {label}", bg=self.t.role("band").fill or "FFFFFF")

    # ---------------- Ausgabe ----------------
    def save(self, outdir: str, stem: str, png_dpi=600, gray_preview=True):
        os.makedirs(outdir, exist_ok=True)
        pptx_path = os.path.join(outdir, stem + ".pptx")
        self.prs.save(pptx_path)
        # PDF über LibreOffice
        tmp = os.path.join(outdir, "_tmp_" + stem)
        os.makedirs(tmp, exist_ok=True)
        pptx_to_pdf(pptx_path, tmp)
        pdf_path = os.path.join(outdir, stem + ".pdf")
        shutil.move(os.path.join(tmp, stem + ".pdf"), pdf_path)
        shutil.rmtree(tmp, ignore_errors=True)
        png_base = os.path.join(outdir, stem)
        subprocess.run(["pdftoppm", "-r", str(png_dpi), "-png", "-singlefile", pdf_path, png_base], check=True)
        if gray_preview:
            gdir = os.path.join(outdir, "_graustufen")
            os.makedirs(gdir, exist_ok=True)
            subprocess.run(["pdftoppm", "-r", "300", "-gray", "-png", "-singlefile", pdf_path,
                            os.path.join(gdir, stem + "_grau")], check=True)
        return pptx_path, pdf_path, png_base + ".png"

    def report(self):
        msgs = list(self.warnings)
        seen = set()
        for lbl, fg, bg in self.text_checks:
            if (fg, bg) in seen:
                continue
            seen.add((fg, bg))
            c = contrast(fg, bg)
            if c < 4.5:
                msgs.append(f"[Kontrast] {fg} auf {bg}: {c:.2f}:1 (< 4,5:1) z. B. bei {lbl}")
        return msgs


# --------------------------------------------------------------------------
# Geometrie-Hilfen
# --------------------------------------------------------------------------
def arc_pts(cx, cy, r, a0, a1, step=2.0):
    """Punkte auf Kreisbogen; Winkel in Grad, mathematisch (0 = rechts, 90 = oben), Bildschirm-y nach unten."""
    n = max(2, int(abs(a1 - a0) / step) + 1)
    return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / (n - 1))),
             cy - r * math.sin(math.radians(a0 + (a1 - a0) * i / (n - 1)))) for i in range(n)]


def polar(cx, cy, r, a):
    return (cx + r * math.cos(math.radians(a)), cy - r * math.sin(math.radians(a)))
