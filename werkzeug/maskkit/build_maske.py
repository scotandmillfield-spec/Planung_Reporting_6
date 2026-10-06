#!/usr/bin/env python3
"""Bündelt eine Eingabemaske zu einer eigenständigen HTML-Datei und erzeugt PNG, PDF und PPTX aus derselben Druckfassung.

Aufruf:
    python3 build_maske.py --maske praemissen.html [--skript praemissen.js] --titel "Planungsbrief" \
        --out out/Exemplarische_Praemissenplanung [--kein-pptx]

Eingabe: HTML-Fragment mit <form class="mk-maske"> (darf eigenes <style> enthalten) und optional ein Skript,
das MK.init({titel}) aufruft, die Logik verdrahtet und mit MK.fertig() endet.
Ausgabe: <out>.html (Lernplattform), <out>.png (600 dpi), <out>.pdf (Vektor, Originalgröße), <out>.pptx (editierbar),
_graustufen/<name>_grau.png und _kontrolle/<name>_pptx.png (PowerPoint gerendert neben dem PNG).
Maßstab fest: 727 px = 110 mm (14 px = 6 pt), also 1 px = 0,1513 mm; eine 1024 px breite Maske wird 155 mm breit.
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile

from PIL import Image, ImageChops, ImageDraw, ImageFont

HIER = os.path.dirname(os.path.abspath(__file__))
MM_PRO_PX = 110.0 / 727.0
DPI = 600
PT_PRO_PX = MM_PRO_PX / (25.4 / 72.0)
EMU_PRO_PX = MM_PRO_PX * 36000.0
CSS_MM_PRO_PX = 25.4 / 96.0          # Chromium: 1 CSS-px = 1/96 Zoll
SOFFICE_WRAPPER = "/mnt/skills/public/pptx/scripts/office/soffice.py"


def lesen(p):
    with open(p, encoding="utf-8") as fh:
        return fh.read()


def buendeln(maske, skript, titel, out_html):
    css = lesen(os.path.join(HIER, "maskkit.css"))
    kit = lesen(os.path.join(HIER, "maskkit.js"))
    frag = lesen(maske)
    js = lesen(skript) if skript else f"MK.init({{ titel: {json.dumps(titel, ensure_ascii=False)} }});\nMK.fertig();"
    sicher = lambda s: s.replace("</script", "<\\/script")
    html = f"""<!DOCTYPE html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{titel}</title>
<style>
{css}
</style>
</head>
<body>
{frag}
<script>
{sicher(kit)}
</script>
<script>
{sicher(js)}
</script>
</body>
</html>
"""
    os.makedirs(os.path.dirname(os.path.abspath(out_html)), exist_ok=True)
    with open(out_html, "w", encoding="utf-8") as fh:
        fh.write(html)
    return out_html


def rendern(html, png, pdf):
    """Druckfassung: PNG (600 dpi), PDF (Vektor) und Exportdaten für PowerPoint."""
    from playwright.sync_api import sync_playwright
    url = "file://" + os.path.abspath(html) + "?print"
    fehler = []
    with sync_playwright() as p:
        b = p.chromium.launch()
        # Breite zuerst bei 1:1 ermitteln
        s0 = b.new_page(viewport={"width": 1200, "height": 900})
        s0.on("pageerror", lambda e: fehler.append(str(e)))
        s0.on("console", lambda m: fehler.append(m.text) if m.type == "error" else None)
        s0.goto(url)
        s0.wait_for_timeout(200)
        s0.evaluate("() => MK.fertig()")
        bw, bh = s0.evaluate("() => { const r = document.querySelector('.mk-maske').getBoundingClientRect(); return [r.width, r.height]; }")
        export = s0.evaluate("() => MK.exportieren()")
        probleme = s0.evaluate("() => MK.pruefen()")
        s0.close()
        # PNG mit 600 dpi
        skala = bw * MM_PRO_PX / 25.4 * DPI / bw
        s1 = b.new_page(viewport={"width": int(bw) + 1, "height": int(bh) + 1}, device_scale_factor=skala)
        s1.goto(url)
        s1.wait_for_timeout(200)
        s1.evaluate("() => MK.fertig()")
        s1.locator(".mk-maske").screenshot(path=png)
        s1.close()
        # PDF in Originalgröße (Text bleibt Text)
        s2 = b.new_page(viewport={"width": int(bw) + 1, "height": int(bh) + 1})
        s2.goto(url)
        s2.wait_for_timeout(200)
        s2.evaluate("() => MK.fertig()")
        s2.add_style_tag(content="@page { margin: 0 } html, body { width: %dpx; height: %dpx; overflow: hidden; }" % (round(bw), round(bh)))
        w_mm, h_mm = bw * MM_PRO_PX, bh * MM_PRO_PX
        s2.pdf(path=pdf, width=f"{w_mm:.3f}mm", height=f"{h_mm + 0.05:.3f}mm", print_background=True,
               scale=MM_PRO_PX / CSS_MM_PRO_PX, margin={"top": "0", "right": "0", "bottom": "0", "left": "0"}, page_ranges="1")
        s2.close()
        b.close()
    im = Image.open(png).convert("RGB")
    im.save(png, dpi=(DPI, DPI))
    gdir = os.path.join(os.path.dirname(os.path.abspath(png)), "_graustufen")
    os.makedirs(gdir, exist_ok=True)
    im.convert("L").resize((im.width // 2, im.height // 2), Image.LANCZOS).save(
        os.path.join(gdir, os.path.basename(png).replace(".png", "_grau.png")))
    return export, fehler, probleme


# ---------------------------------------------------------------- PowerPoint
def pptx_schreiben(export, pfad):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    E = lambda v: Emu(int(round(v * EMU_PRO_PX)))
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(export["breite"]), E(export["hoehe"])
    folie = prs.slides.add_slide(prs.slide_layouts[6])

    def linie_setzen(ln, farbe, lw, dash=False):
        if farbe and lw:
            ln.color.rgb = RGBColor.from_string(farbe)
            ln.width = Pt(lw * PT_PRO_PX)
            if dash:
                ln.dash_style = MSO_LINE.DASH
        else:
            ln.fill.background()

    def ohne_stil(sh):
        # Vorlagenstil entfernen (sonst Schatten und Designfarben aus dem Theme); Linie und Füllung sind explizit gesetzt
        st = sh._element.find(qn("p:style"))
        if st is not None:
            sh._element.remove(st)

    def flaeche(sh, fill):
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            sh.fill.background()

    for f in export["formen"]:
        t = f["t"]
        if t in ("rect", "oval"):
            if f["w"] <= 0 or f["h"] <= 0:
                continue
            art = MSO_SHAPE.OVAL if t == "oval" else (MSO_SHAPE.ROUNDED_RECTANGLE if f.get("r") else MSO_SHAPE.RECTANGLE)
            sh = folie.shapes.add_shape(art, E(f["x"]), E(f["y"]), E(f["w"]), E(f["h"]))
            if art == MSO_SHAPE.ROUNDED_RECTANGLE:
                sh.adjustments[0] = min(0.5, f["r"] / max(1e-6, min(f["w"], f["h"])))
            flaeche(sh, f.get("fill"))
            linie_setzen(sh.line, f.get("line"), f.get("lw", 0), f.get("dash"))
            ohne_stil(sh)
            sh.name = f.get("name", "Form")[:60]
        elif t == "linie":
            if not f.get("farbe"):
                continue
            sh = folie.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(f["x1"]), E(f["y1"]), E(f["x2"]), E(f["y2"]))
            linie_setzen(sh.line, f["farbe"], f["lw"], f.get("dash"))
            ohne_stil(sh)
            sh.name = "Linie"
        elif t == "poly":
            pts = [(int(round(x * EMU_PRO_PX)), int(round(y * EMU_PRO_PX))) for x, y in f["pts"]]
            fb = folie.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
            fb.add_line_segments(pts[1:], close=bool(f.get("zu")))
            sh = fb.convert_to_shape()
            flaeche(sh, f.get("fill"))
            linie_setzen(sh.line, f.get("line"), f.get("lw", 0))
            ohne_stil(sh)
            sh.name = "Form"
        elif t == "text":
            z = f["zeilen"]
            ausr = f.get("ausr", "l")
            x0 = min(q["x"] for q in z)
            x1 = max(q["x"] + q["w"] for q in z)
            y0, y1 = z[0]["y"], z[-1]["y"] + z[-1]["h"]
            reserve = 6.0                                   # Spielraum gegen Rundung, kein Umbruch (wrap=none)
            if ausr == "r":
                x0 -= reserve
            elif ausr == "c":
                x0, x1 = x0 - reserve / 2, x1 + reserve / 2
            else:
                x1 += reserve
            abstand = (z[1]["y"] - z[0]["y"]) if len(z) > 1 else None
            # einzeilig: LibreOffice/PowerPoint setzen die Grundlinie knapp 1 px tiefer als Chromium
            tb = folie.shapes.add_textbox(E(x0), E(y0 - (0.84 if len(z) == 1 else 0.0)), E(x1 - x0), E(max(1.0, y1 - y0)))
            tb.name = ("Text " + f.get("name", ""))[:60]
            tf = tb.text_frame
            tf.word_wrap = False
            tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for i, q in enumerate(z):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = {"r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}.get(ausr, PP_ALIGN.LEFT)
                if abstand:
                    p.line_spacing = Pt(abstand * PT_PRO_PX)
                r = p.add_run()
                r.text = q["text"]
                ft = r.font
                ft.name = "Arial"
                ft.size = Pt(round(f["groesse"] * PT_PRO_PX * 100) / 100)
                ft.bold = bool(f.get("fett"))
                ft.italic = bool(f.get("kursiv"))
                ft.underline = bool(f.get("unter"))
                ft.color.rgb = RGBColor.from_string(f["farbe"])
                rpr = r._r.get_or_add_rPr()
                rpr.set("lang", "de-DE")
                if f.get("sperr"):
                    rpr.set("spc", str(int(round(f["sperr"] * PT_PRO_PX * 100))))
            if abstand:
                # Feste Zeilenhöhe: erste Grundlinie wie im Browser (Ausgleich der halben Zusatzhöhe)
                einfach = f["groesse"] * 1.149
                tb.top = E(y0 - max(0.0, abstand - einfach) * 0.5 + 0.0)
                tb.height = E(abstand * len(z))
    prs.save(pfad)


def pptx_zu_png(pptx, png_ziel, breite_px):
    """PowerPoint über LibreOffice rendern (Kontrolle der Editierfassung)."""
    with tempfile.TemporaryDirectory(prefix="mk_") as tmp:
        if os.path.exists(SOFFICE_WRAPPER):
            cmd = ["python3", SOFFICE_WRAPPER, "--headless", "--convert-to", "pdf", "--outdir", tmp, pptx]
            env = None
        else:
            prof = os.path.join(tmp, "profil")
            cmd = ["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, pptx]
            env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
        subprocess.run(cmd, check=True, capture_output=True, timeout=240, env=env)
        pdf = os.path.join(tmp, os.path.splitext(os.path.basename(pptx))[0] + ".pdf")
        subprocess.run(["pdftoppm", "-r", "300", "-png", "-singlefile", pdf, os.path.join(tmp, "s")], check=True, timeout=120)
        im = Image.open(os.path.join(tmp, "s.png")).convert("RGB")
        im = im.resize((breite_px, round(im.height * breite_px / im.width)), Image.LANCZOS)
        im.save(png_ziel)
        return im


def kontrolle(png, pptx, ziel):
    """PNG (Browser) und PPTX (LibreOffice) nebeneinander und als Differenz; liefert die mittlere Abweichung in %."""
    a = Image.open(png).convert("RGB")
    B = 1300
    a = a.resize((B, round(a.height * B / a.width)), Image.LANCZOS)
    tmp = ziel.replace(".png", "_roh.png")
    b = pptx_zu_png(pptx, tmp, B)
    os.remove(tmp)
    h = max(a.height, b.height)
    b2 = Image.new("RGB", (B, h), "white"); b2.paste(b, (0, 0))
    a2 = Image.new("RGB", (B, h), "white"); a2.paste(a, (0, 0))
    # Abweichung auf weichgezeichneter, verkleinerter Fassung: Kantenglättung und Subpixel zählen nicht, Versatz schon
    from PIL import ImageFilter
    klein = lambda im: im.convert("L").resize((B // 2, h // 2), Image.LANCZOS).filter(ImageFilter.GaussianBlur(1.2))
    diff = ImageChops.difference(klein(a2), klein(b2))
    wert = sum(i * n for i, n in enumerate(diff.histogram())) / ((B // 2) * (h // 2)) / 255 * 100
    # Überlagerung: deckungsgleich schwarz, nur im Browser rot, nur in PowerPoint türkis
    ueber = Image.merge("RGB", (b2.convert("L"), a2.convert("L"), a2.convert("L")))
    blatt = Image.new("RGB", (B * 3 + 80, h + 60), "white")
    d = ImageDraw.Draw(blatt)
    try:
        ft = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf", 26)
    except OSError:
        ft = ImageFont.load_default()
    for i, (titel, bild) in enumerate([("Browser (PNG/PDF)", a2), ("PowerPoint (LibreOffice)", b2), ("Überlagerung: rot nur Browser, türkis nur PowerPoint", ueber)]):
        d.text((i * (B + 40), 10), titel, fill="black", font=ft)
        blatt.paste(bild, (i * (B + 40), 50))
    blatt.save(ziel)
    return wert


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--maske", required=True, help="HTML-Fragment mit <form class='mk-maske'>")
    ap.add_argument("--skript", default=None, help="Masken-Skript (optional)")
    ap.add_argument("--titel", required=True)
    ap.add_argument("--out", required=True, help="Pfad ohne Endung")
    ap.add_argument("--kein-pptx", action="store_true")
    a = ap.parse_args()
    html = buendeln(a.maske, a.skript, a.titel, a.out + ".html")
    print("geschrieben:", html, f"({os.path.getsize(html) // 1024} KB)")
    png, pdf, pptx = a.out + ".png", a.out + ".pdf", a.out + ".pptx"
    export, fehler, probleme = rendern(html, png, pdf)
    g = Image.open(png).size
    print("geschrieben:", png, g, "= %.1f x %.1f mm" % (g[0] / DPI * 25.4, g[1] / DPI * 25.4))
    seiten = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    groesse = [z for z in seiten.splitlines() if z.startswith(("Pages", "Page size"))]
    print("geschrieben:", pdf, " | ".join(" ".join(z.split()) for z in groesse))
    if "Pages:          1" not in seiten and "Pages: 1" not in " ".join(seiten.split()):
        probleme.append("PDF hat mehr als eine Seite")
    with open(a.out + "_export.json", "w", encoding="utf-8") as fh:
        json.dump(export, fh, ensure_ascii=False)
    if not a.kein_pptx:
        pptx_schreiben(export, pptx)
        from pptx import Presentation
        Presentation(pptx)                                    # lässt sich wieder öffnen
        n = len(export["formen"])
        print("geschrieben:", pptx, f"({n} Formen, alle Texte editierbar)")
        kdir = os.path.join(os.path.dirname(os.path.abspath(png)), "_kontrolle")
        os.makedirs(kdir, exist_ok=True)
        ziel = os.path.join(kdir, os.path.basename(a.out) + "_pptx.png")
        try:
            wert = kontrolle(png, pptx, ziel)
            print(f"PowerPoint-Kontrolle: {ziel} (mittlere Abweichung {wert:.2f} %)")
            if wert > 1.2:                                    # 1 px Versatz ergäbe rund 1,8 %
                probleme.append(f"PowerPoint weicht sichtbar vom PNG ab ({wert:.2f} %) – Kontrollbild ansehen")
        except Exception as ex:                               # LibreOffice fehlt: nur Hinweis
            print("PowerPoint-Kontrolle übersprungen:", ex)
    os.remove(a.out + "_export.json")
    probleme = [f"JS-Fehler: {f}" for f in fehler] + probleme
    if probleme:
        print("\nPROBLEME:")
        for p in sorted(set(probleme)):
            print("  ", p)
        sys.exit(1)
    print("Keine Layoutprobleme.")


if __name__ == "__main__":
    main()
